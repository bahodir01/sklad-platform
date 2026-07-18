"""Генерация файлов отчётов: Excel (openpyxl) и PDF (WeasyPrint, §8.1.1/§8.2).

════════════════════════════════════════════════════════════════════════════
ПОЧЕМУ ЭТО ОТДЕЛЬНЫЙ МОДУЛЬ, А НЕ КОД В РОУТЕРЕ (задание этапа 6, §8 архитектуры)
────────────────────────────────────────────────────────────────────────────
По-хорошему тяжёлый экспорт (десятки тысяч строк истории/ДДС) идёт фоновой
задачей Celery (`export_report`, архитектура §8), чтобы не блокировать
uvicorn-воркер. Celery/Redis в этой среде не подняты, поэтому экспорт выполнен
СИНХРОННО, но так, чтобы вынос в фон был тривиальным:

  * генерация — ЧИСТЫЕ функции `table_to_xlsx()` / `table_to_pdf()` над готовой
    структурой `ReportTable` (заголовок + колонки + строки + итоги). Они не
    знают ни про HTTP, ни про сессию БД, ни про FastAPI;
  * сбор данных — `ReportsService.build_*_table()` (service.py): session + фильтры
    → `ReportTable`. Тоже без HTTP.

Чтобы перенести в Celery, задача делает ровно то же, что сейчас делает роутер:
    table = await service.build_balances_table(...)   # сбор данных
    content = table_to_xlsx(table)                     # генерация
и кладёт `content` в MinIO, отдавая клиенту ссылку. Роутер при этом меняется на
«поставить задачу + вернуть 202 + polling», сами функции ниже не трогаются.
════════════════════════════════════════════════════════════════════════════

WeasyPrint импортируется ЛЕНИВО (как в shared/pdf.py): при отсутствии нативных
pango/cairo `table_to_pdf` возвращает готовый к печати HTML вместо PDF —
эндпоинт не притворяется, что выдал PDF, и не падает.
"""

import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import Any

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
PDF_MEDIA_TYPE = "application/pdf"
HTML_MEDIA_TYPE = "text/html; charset=utf-8"


@dataclass(frozen=True)
class Column:
    key: str  # ключ в dict-строке
    header: str  # человекочитаемый заголовок колонки
    numeric: bool = False  # выравнивание/формат числа в Excel


@dataclass
class ReportTable:
    """Формат-независимое представление отчёта. Из него рождаются и xlsx, и pdf,
    и (в service) json. Одна структура — один источник правды для всех форматов."""

    title: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    # Итоговые строки «метка → значение» (балансы касс, суммы прихода/расхода ДДС).
    summary: list[tuple[str, str]] = field(default_factory=list)
    slug: str = "report"  # основа имени файла


@dataclass(frozen=True)
class ExportResult:
    content: bytes
    media_type: str
    filename: str


def _cell(value: Any) -> Any:
    """Значение ячейки Excel: Decimal → float для нативного числового формата,
    date/datetime оставляем как есть (openpyxl отформатирует), None → пусто."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, dt.datetime):
        # openpyxl не умеет tz-aware datetime — снимаем таймзону.
        return value.replace(tzinfo=None)
    return value


def table_to_xlsx(table: ReportTable) -> bytes:
    """ReportTable → .xlsx (openpyxl). Чистая функция (Celery-ready)."""
    from openpyxl import Workbook  # noqa: PLC0415 — импорт рядом с использованием
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = table.title[:31] or "Отчёт"  # лимит Excel на имя листа — 31 символ

    bold = Font(bold=True)
    row_idx = 1
    ws.cell(row=row_idx, column=1, value=table.title).font = bold
    row_idx += 2

    header_row = row_idx
    for col_idx, col in enumerate(table.columns, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=col.header)
        c.font = bold
    row_idx += 1

    for data_row in table.rows:
        for col_idx, col in enumerate(table.columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_cell(data_row.get(col.key)))
        row_idx += 1

    if table.summary:
        row_idx += 1
        for label, value in table.summary:
            ws.cell(row=row_idx, column=1, value=label).font = bold
            ws.cell(row=row_idx, column=2, value=value)
            row_idx += 1

    # Ширина колонок по заголовку (простая эвристика, читаемо без ручной подгонки).
    for col_idx, col in enumerate(table.columns, start=1):
        letter = ws.cell(row=header_row, column=col_idx).column_letter
        ws.column_dimensions[letter].width = max(12, min(40, len(col.header) + 4))
    ws.cell(row=header_row, column=1).alignment = Alignment(horizontal="left")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_html(table: ReportTable) -> str:
    """Готовый к печати HTML отчёта (Jinja2-шаблон templates/pdf/report.html)."""
    from app.shared.pdf import render_template  # noqa: PLC0415

    return render_template(
        "report.html",
        {
            "title": table.title,
            "columns": table.columns,
            "rows": table.rows,
            "summary": table.summary,
            "generated_at": dt.datetime.now().strftime("%d.%m.%Y %H:%M"),
        },
    )


def table_to_pdf(table: ReportTable) -> tuple[bytes, str]:
    """ReportTable → (bytes, media_type). PDF через WeasyPrint; при отсутствии
    нативных зависимостей — HTML-фолбэк (тот же приём, что во всей системе).
    Чистая функция (Celery-ready)."""
    from app.shared.pdf import html_to_pdf  # noqa: PLC0415

    html = _render_html(table)
    pdf = html_to_pdf(html)
    if pdf is not None:
        return pdf, PDF_MEDIA_TYPE
    return html.encode("utf-8"), HTML_MEDIA_TYPE


def export_table(table: ReportTable, fmt: str) -> ExportResult:
    """Единая точка генерации файла отчёта из ReportTable (fmt = xlsx|pdf).

    Это и есть тело будущей Celery-задачи `export_report`: на вход готовая
    таблица, на выход байты + media_type + имя файла. HTTP тут нет."""
    stamp = dt.date.today().strftime("%Y%m%d")
    if fmt == "xlsx":
        return ExportResult(
            content=table_to_xlsx(table),
            media_type=XLSX_MEDIA_TYPE,
            filename=f"{table.slug}_{stamp}.xlsx",
        )
    if fmt == "pdf":
        content, media_type = table_to_pdf(table)
        # Если WeasyPrint недоступен — отдаём .html, честно указывая расширение.
        ext = "pdf" if media_type == PDF_MEDIA_TYPE else "html"
        return ExportResult(
            content=content,
            media_type=media_type,
            filename=f"{table.slug}_{stamp}.{ext}",
        )
    raise ValueError(f"Неизвестный формат экспорта: {fmt!r}")
